"""
Tests for app.services.parser and app.services.exporter.
"""

import io
import sys
import os
import pytest

# Make sure the repo root is on the path so imports work without installing
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from app.services.parser import parse_syllabus_bytes, CaseEntry
from app.services.exporter import to_csv, to_xlsx, entries_to_dataframe, COLUMNS

# ── sample syllabi ────────────────────────────────────────────────────────────

SIMPLE_SYLLABUS = """\
1. Contract Law
   Smith v Jones [1990]
1.1 Offer and Acceptance
   Adams v Lindsell [1818]
1.1.1 Communication of Offer
   Byrne v Van Tienhoven [1880]
1.1.1.1 Postal Rule
   Household Fire Insurance v Grant [1879]
2. Tort Law
   Donoghue v Stevenson [1932]
2.1 Negligence
   Caparo Industries v Dickman [1990]
"""

MULTI_CASE_TOPIC = """\
1. Criminal Law
   R v Brown [1994] and R v Smith [2000]
1.1 Murder
   R v Woollin [1998]
"""

NO_CASES_SYLLABUS = """\
1. Introduction
   This section covers general principles.
1.1 Overview
   No specific cases are cited here.
"""


# ── parser tests ──────────────────────────────────────────────────────────────

def test_parse_simple_txt():
    entries = parse_syllabus_bytes(SIMPLE_SYLLABUS.encode(), "syllabus.txt")
    assert len(entries) >= 6, "Expected at least 6 case entries"


def test_parse_topic_assignment():
    entries = parse_syllabus_bytes(SIMPLE_SYLLABUS.encode(), "syllabus.txt")
    # Smith v Jones should be under topic "Contract Law"
    smith = next((e for e in entries if "Smith" in e.case), None)
    assert smith is not None
    assert smith.topic == "Contract Law"
    assert smith.subtopic == ""


def test_parse_subtopic_assignment():
    entries = parse_syllabus_bytes(SIMPLE_SYLLABUS.encode(), "syllabus.txt")
    adams = next((e for e in entries if "Adams" in e.case), None)
    assert adams is not None
    assert adams.topic == "Contract Law"
    assert adams.subtopic == "Offer and Acceptance"


def test_parse_sub_subtopic_assignment():
    entries = parse_syllabus_bytes(SIMPLE_SYLLABUS.encode(), "syllabus.txt")
    byrne = next((e for e in entries if "Byrne" in e.case), None)
    assert byrne is not None
    assert byrne.sub_subtopic == "Communication of Offer"


def test_parse_further_division():
    entries = parse_syllabus_bytes(SIMPLE_SYLLABUS.encode(), "syllabus.txt")
    household = next((e for e in entries if "Household" in e.case), None)
    assert household is not None
    assert household.further_division == "Postal Rule"


def test_parse_no_cases():
    entries = parse_syllabus_bytes(NO_CASES_SYLLABUS.encode(), "syllabus.txt")
    assert entries == []


def test_parse_multi_case_on_one_line():
    entries = parse_syllabus_bytes(MULTI_CASE_TOPIC.encode(), "syllabus.txt")
    criminal_cases = [e for e in entries if e.topic == "Criminal Law" and e.subtopic == ""]
    assert len(criminal_cases) >= 2


def test_deduplication():
    # Same case twice on the same topic → should appear only once
    syllabus = "1. Contract Law\n   Smith v Jones [1990]\n   Smith v Jones [1990]\n"
    entries = parse_syllabus_bytes(syllabus.encode(), "syllabus.txt")
    names = [e.case for e in entries]
    assert names.count("Smith v Jones [1990]") <= 1 or len([e for e in entries if "Smith" in e.case]) == 1


# ── exporter tests ────────────────────────────────────────────────────────────

def _sample_entries():
    return [
        CaseEntry("Smith v Jones [1990]", "Contract Law", "Offer", "", ""),
        CaseEntry("Donoghue v Stevenson [1932]", "Tort Law", "Negligence", "", ""),
    ]


def test_dataframe_columns():
    df = entries_to_dataframe(_sample_entries())
    assert list(df.columns) == COLUMNS


def test_dataframe_row_count():
    df = entries_to_dataframe(_sample_entries())
    assert len(df) == 2


def test_csv_output():
    data = to_csv(_sample_entries())
    text = data.decode("utf-8")
    assert "Case,Topic,Subtopic,Sub-subtopic,Further Division" in text
    assert "Smith v Jones" in text


def test_xlsx_output():
    import openpyxl
    data = to_xlsx(_sample_entries())
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    assert headers == COLUMNS


def test_empty_entries_csv():
    data = to_csv([])
    text = data.decode("utf-8")
    assert "Case" in text  # headers still present


def test_empty_entries_xlsx():
    import openpyxl
    data = to_xlsx([])
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    assert headers == COLUMNS
