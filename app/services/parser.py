"""
Syllabus parser – extracts the hierarchical topic structure and cases
from uploaded syllabus files (PDF, DOCX, TXT).

Hierarchy levels detected
─────────────────────────
Level 1  Topic              e.g.  "1.  Contract Law"
Level 2  Subtopic           e.g.  "1.1  Offer and Acceptance"
Level 3  Sub-subtopic       e.g.  "1.1.1  Communication of Offer"
Level 4  Further Division   e.g.  "1.1.1.1  Postal Rule"

Cases are typically written as
  • "Smith v Jones [1990]"
  • "Re ABC Ltd"
  • "XYZ v ABC"
and appear inside the text at any hierarchy level.
"""

from __future__ import annotations

import io
import os
import re
from dataclasses import dataclass, field
from typing import List, Optional

# ── optional heavy imports (graceful fallback) ──────────────────────────────
try:
    import pdfplumber  # type: ignore
    _HAS_PDF = True
except ImportError:
    _HAS_PDF = False

try:
    from docx import Document as DocxDocument  # type: ignore
    _HAS_DOCX = True
except ImportError:
    _HAS_DOCX = False


# ── regex helpers ────────────────────────────────────────────────────────────

# Numbered-heading patterns  (trailing ., ) or – are optional)
_RE_TOPIC          = re.compile(r"^\s*(\d+)\s*[.\)–\-]?\s+(.+)", re.IGNORECASE)
_RE_SUBTOPIC       = re.compile(r"^\s*(\d+\.\d+)\s*[.\)–\-]?\s+(.+)", re.IGNORECASE)
_RE_SUBSUBTOPIC    = re.compile(r"^\s*(\d+\.\d+\.\d+)\s*[.\)–\-]?\s+(.+)", re.IGNORECASE)
_RE_FURTHERDIV     = re.compile(r"^\s*(\d+\.\d+\.\d+\.\d+)\s*[.\)–\-]?\s+(.+)", re.IGNORECASE)

# Roman-numeral / lettered sub-items (i. ii. iii. / a. b. c.)
_RE_ROMAN          = re.compile(
    r"^\s*((?:x{0,3})(?:ix|iv|v?i{0,3}))\s*[.\)]\s+(.+)", re.IGNORECASE
)
_RE_ALPHA          = re.compile(r"^\s*([a-zA-Z])\s*[.\)]\s+(.+)")

# Case citation patterns
# Handles:  "Smith v Jones", "R v Brown", "Re Smith", "In re ABC", "Ex parte XYZ"
# Single-letter names (R, A, B…) are common in criminal-law citations.
_RE_CASE = re.compile(
    r"\b(?:"
    r"[A-Z][a-zA-Z]*(?: [A-Z][a-zA-Z]*)* v\.? [A-Z][a-zA-Z]+(?: [A-Z][a-zA-Z]+)*"
    r"|Re [A-Z][a-zA-Z]+(?: [A-Z][a-zA-Z]+)*"
    r"|In re [A-Z][a-zA-Z]+(?: [A-Z][a-zA-Z]+)*"
    r"|Ex parte [A-Z][a-zA-Z]+(?: [A-Z][a-zA-Z]+)*"
    r")(?:\s*\[?\d{4}\]?)?",
)

# Bullet / dash list item
_RE_BULLET = re.compile(r"^\s*[-•*]\s+(.+)")


# ── data structures ──────────────────────────────────────────────────────────

@dataclass
class CaseEntry:
    case: str
    topic: str = ""
    subtopic: str = ""
    sub_subtopic: str = ""
    further_division: str = ""

    def as_dict(self) -> dict:
        return {
            "Case": self.case,
            "Topic": self.topic,
            "Subtopic": self.subtopic,
            "Sub-subtopic": self.sub_subtopic,
            "Further Division": self.further_division,
        }


@dataclass
class _ParseState:
    topic: str = ""
    subtopic: str = ""
    sub_subtopic: str = ""
    further_division: str = ""


# ── public API ───────────────────────────────────────────────────────────────

def parse_syllabus(filepath: str) -> List[CaseEntry]:
    """
    Parse *filepath* and return a list of :class:`CaseEntry` objects.

    The function auto-detects the file format by extension.
    """
    ext = os.path.splitext(filepath)[1].lower()
    if ext == ".pdf":
        lines = _lines_from_pdf(filepath)
    elif ext in (".docx", ".doc"):
        lines = _lines_from_docx(filepath)
    else:
        lines = _lines_from_txt(filepath)

    return _extract_entries(lines)


def parse_syllabus_bytes(data: bytes, filename: str) -> List[CaseEntry]:
    """Like :func:`parse_syllabus` but works on an in-memory *bytes* object."""
    ext = os.path.splitext(filename)[1].lower()
    if ext == ".pdf":
        lines = _lines_from_pdf_bytes(data)
    elif ext in (".docx", ".doc"):
        lines = _lines_from_docx_bytes(data)
    else:
        lines = _lines_from_txt_bytes(data)

    return _extract_entries(lines)


# ── text extraction ──────────────────────────────────────────────────────────

def _lines_from_txt(filepath: str) -> List[str]:
    with open(filepath, encoding="utf-8", errors="replace") as fh:
        return fh.read().splitlines()


def _lines_from_txt_bytes(data: bytes) -> List[str]:
    return data.decode("utf-8", errors="replace").splitlines()


def _lines_from_pdf(filepath: str) -> List[str]:
    if not _HAS_PDF:
        raise RuntimeError("pdfplumber is not installed")
    with pdfplumber.open(filepath) as pdf:
        return _pdf_pages_to_lines(pdf)


def _lines_from_pdf_bytes(data: bytes) -> List[str]:
    if not _HAS_PDF:
        raise RuntimeError("pdfplumber is not installed")
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        return _pdf_pages_to_lines(pdf)


def _pdf_pages_to_lines(pdf) -> List[str]:
    lines: List[str] = []
    for page in pdf.pages:
        text = page.extract_text() or ""
        lines.extend(text.splitlines())
    return lines


def _lines_from_docx(filepath: str) -> List[str]:
    if not _HAS_DOCX:
        raise RuntimeError("python-docx is not installed")
    doc = DocxDocument(filepath)
    return [p.text for p in doc.paragraphs]


def _lines_from_docx_bytes(data: bytes) -> List[str]:
    if not _HAS_DOCX:
        raise RuntimeError("python-docx is not installed")
    doc = DocxDocument(io.BytesIO(data))
    return [p.text for p in doc.paragraphs]


# ── core parsing logic ───────────────────────────────────────────────────────

def _extract_entries(lines: List[str]) -> List[CaseEntry]:
    """
    Walk through *lines* and produce :class:`CaseEntry` records.

    Strategy
    ────────
    1.  Keep track of the current heading context (_ParseState).
    2.  On every line look for:
        a. A deeper heading (updates context).
        b. Case citations (creates a CaseEntry under current context).
    3.  The same line may both update context AND contain cases.
    """
    state = _ParseState()
    entries: List[CaseEntry] = []

    for raw in lines:
        line = raw.strip()
        if not line:
            continue

        # ── detect heading level ────────────────────────────────────────────
        m4 = _RE_FURTHERDIV.match(line)
        m3 = _RE_SUBSUBTOPIC.match(line)
        m2 = _RE_SUBTOPIC.match(line)
        m1 = _RE_TOPIC.match(line)

        # Use the most-specific match first
        if m4:
            state.further_division = m4.group(2).strip()
        elif m3:
            state.sub_subtopic = m3.group(2).strip()
            state.further_division = ""
        elif m2:
            state.subtopic = m2.group(2).strip()
            state.sub_subtopic = ""
            state.further_division = ""
        elif m1:
            state.topic = m1.group(2).strip()
            state.subtopic = ""
            state.sub_subtopic = ""
            state.further_division = ""

        # ── extract cases from the line text ───────────────────────────────
        # Strip bullet / numbering prefix to get the bare text
        bare = line
        for pat in (_RE_BULLET, _RE_ALPHA, _RE_ROMAN,
                    _RE_FURTHERDIV, _RE_SUBSUBTOPIC, _RE_SUBTOPIC, _RE_TOPIC):
            bm = pat.match(bare)
            if bm:
                bare = bm.group(bm.lastindex).strip()
                break

        case_hits = _RE_CASE.findall(bare)
        for hit in case_hits:
            case_name = hit.strip()
            entries.append(
                CaseEntry(
                    case=case_name,
                    topic=state.topic,
                    subtopic=state.subtopic,
                    sub_subtopic=state.sub_subtopic,
                    further_division=state.further_division,
                )
            )

    # ── de-duplicate while preserving order ────────────────────────────────
    seen: set = set()
    unique: List[CaseEntry] = []
    for e in entries:
        key = (e.case, e.topic, e.subtopic, e.sub_subtopic, e.further_division)
        if key not in seen:
            seen.add(key)
            unique.append(e)

    return unique
